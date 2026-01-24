"""Excel 导出测试"""
import io
import pytest
from openpyxl import load_workbook
from app.services.excel_export import ExcelExporter
from app.models.dimensions import (
    FunctionalDimensions,
    TechnicalDimensions,
    PricingDimensions,
    CostCalculationInput,
)
from app.models.enums import RecordingMode, VideoQuality, StorageClass
from app.services.calculator.s3_standard import S3StandardCalculator
from app.services.calculator.comparator import StorageComparator


class TestExcelExporter:
    """Excel 导出器测试"""

    @pytest.fixture
    def sample_input(self):
        """示例输入"""
        return CostCalculationInput(
            functional=FunctionalDimensions(
                device_count=1000,
                recording_mode=RecordingMode.EVENT_TRIGGERED,
                video_quality=VideoQuality.P1080,
                events_per_day=400,
                event_duration_sec=15,
                retention_days=30,
                access_pattern=0.1,
            ),
            technical=TechnicalDimensions(
                storage_class=StorageClass.STANDARD,
            ),
            pricing=PricingDimensions(
                region="ap-northeast-1",
                discount_percent=0.0,
            ),
        )

    @pytest.fixture
    def sample_result(self, sample_input):
        """示例计算结果"""
        calculator = S3StandardCalculator()
        return calculator.calculate(sample_input)

    @pytest.fixture
    def sample_comparison(self, sample_input):
        """示例对比结果"""
        comparator = StorageComparator()
        return comparator.compare(sample_input)

    def test_generate_returns_bytes(self, sample_input, sample_result):
        """测试生成返回字节流"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
        )

        assert isinstance(content, bytes)
        assert len(content) > 0

    def test_generated_file_is_valid_excel(self, sample_input, sample_result):
        """测试生成的文件是有效的 Excel"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
        )

        # 尝试加载为 Excel
        wb = load_workbook(io.BytesIO(content))
        assert wb is not None
        assert len(wb.sheetnames) > 0

    def test_has_summary_sheet(self, sample_input, sample_result):
        """测试有汇总表"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
        )

        wb = load_workbook(io.BytesIO(content))
        assert "成本汇总" in wb.sheetnames

    def test_has_breakdown_sheet(self, sample_input, sample_result):
        """测试有明细表"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
        )

        wb = load_workbook(io.BytesIO(content))
        assert "费用明细" in wb.sheetnames

    def test_has_input_params_sheet(self, sample_input, sample_result):
        """测试有输入参数表"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
        )

        wb = load_workbook(io.BytesIO(content))
        assert "输入参数" in wb.sheetnames

    def test_summary_contains_cost_data(self, sample_input, sample_result):
        """测试汇总表包含成本数据"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
        )

        wb = load_workbook(io.BytesIO(content))
        ws = wb["成本汇总"]

        # 检查是否包含月度总费用
        found_monthly = False
        for row in ws.iter_rows(values_only=True):
            if row and "月度" in str(row[0] or ""):
                found_monthly = True
                break
        assert found_monthly

    def test_with_comparison(self, sample_input, sample_result, sample_comparison):
        """测试带方案对比"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
            comparison=sample_comparison,
        )

        wb = load_workbook(io.BytesIO(content))
        assert "方案对比" in wb.sheetnames

    def test_comparison_has_multiple_options(self, sample_input, sample_result, sample_comparison):
        """测试方案对比包含多个选项"""
        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=sample_input,
            result=sample_result,
            comparison=sample_comparison,
        )

        wb = load_workbook(io.BytesIO(content))
        ws = wb["方案对比"]

        # 检查是否包含多个方案
        options_found = 0
        for row in ws.iter_rows(values_only=True):
            if row and ("Standard" in str(row[0] or "") or "Glacier" in str(row[0] or "")):
                options_found += 1
        assert options_found >= 2


class TestExcelExporterEdgeCases:
    """Excel 导出器边界测试"""

    def test_single_device(self):
        """测试单设备场景"""
        input_data = CostCalculationInput(
            functional=FunctionalDimensions(
                device_count=1,
                recording_mode=RecordingMode.EVENT_TRIGGERED,
            ),
            technical=TechnicalDimensions(),
            pricing=PricingDimensions(),
        )

        calculator = S3StandardCalculator()
        result = calculator.calculate(input_data)

        exporter = ExcelExporter()
        content = exporter.generate(input_data=input_data, result=result)

        assert len(content) > 0

    def test_continuous_recording(self):
        """测试全天候录像"""
        input_data = CostCalculationInput(
            functional=FunctionalDimensions(
                device_count=100,
                recording_mode=RecordingMode.CONTINUOUS,
                retention_days=30,
            ),
            technical=TechnicalDimensions(),
            pricing=PricingDimensions(),
        )

        calculator = S3StandardCalculator()
        result = calculator.calculate(input_data)

        exporter = ExcelExporter()
        content = exporter.generate(input_data=input_data, result=result)

        assert len(content) > 0

    def test_custom_title(self):
        """测试自定义标题"""
        input_data = CostCalculationInput(
            functional=FunctionalDimensions(
                device_count=100,
                recording_mode=RecordingMode.EVENT_TRIGGERED,
            ),
            technical=TechnicalDimensions(),
            pricing=PricingDimensions(),
        )

        calculator = S3StandardCalculator()
        result = calculator.calculate(input_data)

        exporter = ExcelExporter()
        content = exporter.generate(
            input_data=input_data,
            result=result,
            title="自定义评估报告",
        )

        wb = load_workbook(io.BytesIO(content))
        ws = wb["成本汇总"]
        # 第一行应该是标题
        assert "自定义评估报告" in str(ws["A1"].value or "")
